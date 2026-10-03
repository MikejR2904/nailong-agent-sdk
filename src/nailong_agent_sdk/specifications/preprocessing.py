# Copyright (c) 2026 David Michael Indraputra

"""Manifest-driven, source-preserving specification preprocessing.

The output is an ordered document tree with document/source locations. It is
not a lossy summary or a generic RAG index (systems-design framework, pp. 22-30).
"""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import yaml

from .documents import (
    DocumentFormat,
    DocumentNode,
    DocumentNodeKind,
    DocumentTree,
    SourceRef,
    SpecificationCategory,
    SpecificationDocument,
    SpecificationManifest,
    VisionStatus,
)
from .vision import VisionAdapter, VisionProposal


class SpecificationPreprocessor:
    """Safely parse manifest files and source documents beneath one root."""

    def __init__(self, specification_root: Path) -> None:
        self._root = specification_root.resolve()

    def load_manifest(
        self, relative_manifest_path: str = "specification-manifest.yaml"
    ) -> SpecificationManifest:
        path = self._resolve(relative_manifest_path)
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            raise ValueError("Specification manifest must be a mapping.")
        if "documents" not in payload:
            payload = self._legacy_manifest_to_documents(payload)
        return SpecificationManifest.model_validate(payload)

    def process_manifest(self, manifest: SpecificationManifest) -> list[DocumentTree]:
        return [self.process_document(document) for document in manifest.documents]

    def process_document(self, document: SpecificationDocument) -> DocumentTree:
        path = self._resolve(document.path)
        if not path.is_file():
            raise ValueError(f'Specification document "{document.path}" does not exist.')
        raw = path.read_bytes()
        digest = hashlib.sha256(raw).hexdigest()

        def source(location: str) -> SourceRef:
            return SourceRef(
                document_id=document.id,
                relative_path=document.path,
                source_hash=digest,
                format=document.format,
                location=location,
            )

        nodes = self._parse(path, document.format, source)
        return DocumentTree(
            document_id=document.id,
            category=document.category,
            title=document.title,
            format=document.format,
            relative_path=document.path,
            source_hash=digest,
            nodes=self._with_context(nodes),
        )

    async def resolve_images(
        self,
        tree: DocumentTree,
        adapter: VisionAdapter | None = None,
        *,
        confidence_threshold: float = 0.8,
        max_attempts: int = 3,
    ) -> DocumentTree:
        """Resolve image nodes through ``adapter``.

        With no adapter configured, image nodes are marked ``REVIEW_REQUIRED`` so a
        person interprets them; nothing is guessed.
        """

        resolved: list[DocumentNode] = []
        for node in tree.nodes:
            if node.kind is not DocumentNodeKind.IMAGE:
                resolved.append(node)
                continue
            if adapter is None:
                resolved.append(
                    node.model_copy(update={"vision_status": VisionStatus.REVIEW_REQUIRED})
                )
                continue
            proposal: VisionProposal | None = None
            for _ in range(max_attempts):
                candidate = await adapter.extract(node)
                if candidate.confidence >= confidence_threshold and candidate.structure:
                    proposal = candidate
                    break
                proposal = candidate
            if proposal and proposal.confidence >= confidence_threshold and proposal.structure:
                resolved.append(
                    node.model_copy(
                        update={
                            "vision_status": VisionStatus.ACCEPTED,
                            "resolved_structure": proposal.structure,
                        }
                    )
                )
            else:
                resolved.append(
                    node.model_copy(update={"vision_status": VisionStatus.REVIEW_REQUIRED})
                )
        return tree.model_copy(update={"nodes": resolved})

    def persist_tree(self, tree: DocumentTree) -> Path:
        target = self._root / "processed" / f"{tree.document_id}.document.yaml"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            yaml.safe_dump(tree.model_dump(mode="json"), sort_keys=False), encoding="utf-8"
        )
        return target

    def _parse(self, path: Path, format_: DocumentFormat, source) -> list[DocumentNode]:
        if format_ in {
            DocumentFormat.TXT,
            DocumentFormat.MD,
            DocumentFormat.TEX,
            DocumentFormat.SYSTEMRDL,
            DocumentFormat.SDC,
            DocumentFormat.UPF,
            DocumentFormat.PLANTUML,
            DocumentFormat.DOT,
            DocumentFormat.WAVEDROM,
            DocumentFormat.MERMAID,
            DocumentFormat.TIKZ,
        }:
            kind = (
                DocumentNodeKind.DIAGRAM
                if format_
                in {
                    DocumentFormat.PLANTUML,
                    DocumentFormat.DOT,
                    DocumentFormat.WAVEDROM,
                    DocumentFormat.MERMAID,
                    DocumentFormat.TIKZ,
                }
                else DocumentNodeKind.TEXT
            )
            return self._line_nodes(
                path.read_text(encoding="utf-8", errors="replace"), source, kind
            )
        if format_ is DocumentFormat.PDF:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            nodes: list[DocumentNode] = []
            for index, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                nodes.append(
                    DocumentNode(
                        node_id=f"page-{index}-text",
                        kind=DocumentNodeKind.TEXT,
                        source=source(f"page:{index}"),
                        content=text,
                    )
                )
                for image_index, image in enumerate(page.images, start=1):
                    nodes.append(
                        DocumentNode(
                            node_id=f"page-{index}-image-{image_index}",
                            kind=DocumentNodeKind.IMAGE,
                            source=source(f"page:{index};image:{image_index}"),
                            content={"name": image.name},
                            vision_status=VisionStatus.PENDING,
                        )
                    )
            return nodes
        if format_ is DocumentFormat.DOCX:
            from docx import Document

            document = Document(str(path))
            nodes = []
            image_index = 0
            for index, paragraph in enumerate(document.paragraphs, start=1):
                if paragraph.text.strip():
                    nodes.append(
                        DocumentNode(
                            node_id=f"paragraph-{index}",
                            kind=DocumentNodeKind.TEXT,
                            source=source(f"paragraph:{index}"),
                            content=paragraph.text,
                        )
                    )
                drawing_count = sum(
                    child.tag.endswith("}drawing")
                    for child in paragraph._p.iter()  # noqa: SLF001
                )
                for _ in range(drawing_count):
                    image_index += 1
                    nodes.append(
                        DocumentNode(
                            node_id=f"paragraph-{index}-image-{image_index}",
                            kind=DocumentNodeKind.IMAGE,
                            source=source(f"paragraph:{index};image:{image_index}"),
                            content={"anchor": "paragraph-inline-image"},
                            vision_status=VisionStatus.PENDING,
                        )
                    )
            for index, table in enumerate(document.tables, start=1):
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                nodes.append(
                    DocumentNode(
                        node_id=f"table-{index}",
                        kind=DocumentNodeKind.TABLE,
                        source=source(f"table:{index}"),
                        content=rows,
                    )
                )
            return nodes
        if format_ in {DocumentFormat.CSV, DocumentFormat.XLSX}:
            return self._table_nodes(path, format_, source)
        if format_ in {DocumentFormat.YAML, DocumentFormat.JSON}:
            parsed = (
                yaml.safe_load(path.read_text(encoding="utf-8"))
                if format_ is DocumentFormat.YAML
                else json.loads(path.read_text(encoding="utf-8"))
            )
            return [
                DocumentNode(
                    node_id="structured-1",
                    kind=DocumentNodeKind.STRUCTURED,
                    source=source("root"),
                    content=parsed,
                )
            ]
        if format_ in {
            DocumentFormat.XML,
            DocumentFormat.SVG,
            DocumentFormat.DRAWIO,
        }:
            from defusedxml import ElementTree

            tree = ElementTree.parse(path)
            root = tree.getroot()
            return [
                DocumentNode(
                    node_id="xml-root",
                    kind=DocumentNodeKind.STRUCTURED
                    if format_ is DocumentFormat.XML
                    else DocumentNodeKind.DIAGRAM,
                    source=source(f"xml:/{root.tag}"),
                    content={
                        "tag": root.tag,
                        "attributes": dict(root.attrib),
                        "xml": ElementTree.tostring(root, encoding="unicode"),
                    },
                )
            ]
        if format_ is DocumentFormat.VSDX:
            with zipfile.ZipFile(path) as package:
                page_parts = sorted(
                    item
                    for item in package.namelist()
                    if item.startswith("visio/pages/") and item.endswith(".xml")
                )
                if not page_parts:
                    raise ValueError("VSDX document contains no Visio page XML parts.")
                return [
                    DocumentNode(
                        node_id=f"vsdx-{index}",
                        kind=DocumentNodeKind.DIAGRAM,
                        source=source(f"vsdx-part:{part}"),
                        content={
                            "part": part,
                            "xml": package.read(part).decode("utf-8", errors="replace"),
                        },
                    )
                    for index, part in enumerate(page_parts, start=1)
                ]
        if format_ in {DocumentFormat.PNG, DocumentFormat.JPEG}:
            return [
                DocumentNode(
                    node_id="image-1",
                    kind=DocumentNodeKind.IMAGE,
                    source=source("image:1"),
                    content={"filename": path.name},
                    vision_status=VisionStatus.PENDING,
                )
            ]
        raise ValueError(f'Unsupported specification format "{format_.value}".')

    @staticmethod
    def _line_nodes(text: str, source, kind: DocumentNodeKind) -> list[DocumentNode]:
        nodes = []
        for index, line in enumerate(text.splitlines(), start=1):
            if line.strip():
                nodes.append(
                    DocumentNode(
                        node_id=f"line-{index}",
                        kind=kind,
                        source=source(f"line:{index}"),
                        content=line,
                    )
                )
        return nodes or [
            DocumentNode(node_id="line-1", kind=kind, source=source("line:1"), content="")
        ]

    @staticmethod
    def _table_nodes(path: Path, format_: DocumentFormat, source) -> list[DocumentNode]:
        if format_ is DocumentFormat.CSV:
            import csv

            with path.open(newline="", encoding="utf-8") as file:
                return [
                    DocumentNode(
                        node_id="table-1",
                        kind=DocumentNodeKind.TABLE,
                        source=source("rows:1-end"),
                        content=list(csv.reader(file)),
                    )
                ]
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        nodes = []
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            nodes.append(
                DocumentNode(
                    node_id=f"sheet-{sheet.title}",
                    kind=DocumentNodeKind.TABLE,
                    source=source(f"sheet:{sheet.title};range:{sheet.calculate_dimension()}"),
                    content=rows,
                )
            )
        return nodes

    @staticmethod
    def _with_context(nodes: list[DocumentNode]) -> list[DocumentNode]:
        values = []
        text_values = [str(node.content) for node in nodes]
        for index, node in enumerate(nodes):
            surrounding = [
                item
                for item in text_values[max(0, index - 1) : index + 2]
                if item != str(node.content)
            ]
            values.append(node.model_copy(update={"surrounding_text": surrounding[:2]}))
        return values

    def _resolve(self, relative_path: str) -> Path:
        candidate = (self._root / relative_path).resolve()
        try:
            candidate.relative_to(self._root)
        except ValueError as error:
            raise ValueError("Specification path escapes the configured root.") from error
        return candidate

    @staticmethod
    def _legacy_manifest_to_documents(payload: dict[str, Any]) -> dict[str, Any]:
        documents: list[dict[str, Any]] = []
        categories = {
            "functional_spec": SpecificationCategory.FUNCTIONAL,
            "architectural_spec": SpecificationCategory.ARCHITECTURAL,
            "interface_spec": SpecificationCategory.INTERFACE,
            "ppa_spec": SpecificationCategory.PPA,
            "pdk_spec": SpecificationCategory.PDK,
            "verification_spec": SpecificationCategory.VERIFICATION,
            "safety_security_spec": SpecificationCategory.SAFETY_SECURITY,
            "assumptions_dependencies": SpecificationCategory.ASSUMPTIONS,
        }
        for key, category in categories.items():
            section = payload.get(key, {})
            for item in section.get("documents", []):
                documents.append({**item, "category": category.value})
            for interface in section.get("interfaces", []):
                for item in interface.get("documents", []):
                    documents.append({**item, "category": category.value})
        return {"documents": documents}


def keywords_from_task(text: str) -> set[str]:
    return {token.lower() for token in re.findall(r"[A-Za-z][A-Za-z0-9_-]{2,}", text)}
