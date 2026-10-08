import asyncio
import itertools
import json
import random
import time
import zipfile

import pytest
import yaml

from nailong_agent_sdk.specifications.documents import (
    DocumentFormat,
    DocumentNodeKind,
    SpecificationCategory,
    SpecificationDocument,
    SpecificationManifest,
    VisionStatus,
)
from nailong_agent_sdk.specifications.evidence_graph import (
    EvidenceGraph,
    EvidenceNode,
    EvidenceNodeKind,
    EvidenceRelation,
    EvidenceRelationKind,
    EvidenceSelectionPolicy,
    EvidenceSelectionRequest,
    EvidenceSelectionStatus,
    StructuralContextSelector,
)
from nailong_agent_sdk.specifications.gate import (
    Gate1ArtifactStore,
    SpecificationGate,
    classify_version_change,
)
from nailong_agent_sdk.specifications.gate_models import (
    GapSeverity,
    GapType,
    SemanticGapFinding,
    VersionChangeKind,
    VersionMetadata,
)
from nailong_agent_sdk.specifications.preprocessing import (
    SpecificationPreprocessor,
    keywords_from_task,
)
from nailong_agent_sdk.specifications.vision import (
    ScriptedVisionAdapter,
    UnconfiguredVisionAdapter,
    VisionProposal,
)
from tests.support.specs import node, req, spec, sref, tree


def arun(coro, timeout=60):
    async def guarded():
        return await asyncio.wait_for(coro, timeout)

    return asyncio.run(guarded())


def document(path, fmt, doc_id="d1", category=SpecificationCategory.FUNCTIONAL):
    return SpecificationDocument(id=doc_id, title=doc_id, format=fmt, path=path, category=category)


def write(root, name, content, binary=False):
    target = root / name
    target.parent.mkdir(parents=True, exist_ok=True)
    if binary:
        target.write_bytes(content)
    else:
        target.write_text(content, "utf-8")
    return target


def test_manifest_load_new_legacy_and_errors(tmp_path):
    write(
        tmp_path,
        "specification-manifest.yaml",
        yaml.safe_dump(
            {
                "documents": [
                    {
                        "id": "a",
                        "title": "A",
                        "format": "md",
                        "path": "a.md",
                        "category": "functional",
                    }
                ]
            }
        ),
    )
    pre = SpecificationPreprocessor(tmp_path)
    manifest = pre.load_manifest()
    assert manifest.documents[0].id == "a"
    legacy = {
        "functional_spec": {
            "documents": [{"id": "f", "title": "F", "format": "md", "path": "f.md"}]
        },
        "interface_spec": {
            "interfaces": [
                {"documents": [{"id": "i", "title": "I", "format": "md", "path": "i.md"}]}
            ]
        },
    }
    write(tmp_path, "legacy.yaml", yaml.safe_dump(legacy))
    converted = pre.load_manifest("legacy.yaml")
    assert {d.id: d.category.value for d in converted.documents} == {
        "f": "functional",
        "i": "interface",
    }
    write(tmp_path, "list.yaml", "- 1\n- 2\n")
    with pytest.raises(ValueError, match="must be a mapping"):
        pre.load_manifest("list.yaml")
    with pytest.raises(ValueError, match="escapes the configured root"):
        pre.load_manifest("../outside.yaml")
    with pytest.raises(ValueError):
        SpecificationDocument(id="x", title="x", format="md", path="../x.md", category="functional")
    for rooted in ("/abs.md", "\\abs.md", "C:abs.md", "C:/abs.md"):
        with pytest.raises(ValueError, match="must be relative"):
            SpecificationDocument(
                id="x", title="x", format="md", path=rooted, category="functional"
            )
    with pytest.raises(ValueError, match="unique"):
        SpecificationManifest(documents=[document("a.md", "md"), document("b.md", "md")])
    empty = write(tmp_path, "empty.yaml", "")
    with pytest.raises(ValueError, match="must be a mapping"):
        pre.load_manifest("empty.yaml")
    assert empty.exists()


def test_text_nodes_locations_hashes_and_context(tmp_path):
    write(tmp_path, "a.md", "alpha\n\nbeta\ngamma\n")
    pre = SpecificationPreprocessor(tmp_path)
    result = pre.process_document(document("a.md", DocumentFormat.MD))
    assert [n.node_id for n in result.nodes] == ["line-1", "line-3", "line-4"]
    assert [n.source.location for n in result.nodes] == ["line:1", "line:3", "line:4"]
    assert result.nodes[1].surrounding_text == ["alpha", "gamma"]
    import hashlib

    assert result.source_hash == hashlib.sha256((tmp_path / "a.md").read_bytes()).hexdigest()
    assert all(n.source.source_hash == result.source_hash for n in result.nodes)
    write(tmp_path, "empty.md", "")
    only = pre.process_document(document("empty.md", DocumentFormat.MD))
    assert [n.node_id for n in only.nodes] == ["line-1"] and only.nodes[0].content == ""
    write(tmp_path, "dups.md", "same\nsame\nsame\n")
    pre.process_document(document("dups.md", DocumentFormat.MD))


def test_location_numbers_follow_newlines_not_splitlines(tmp_path):
    write(tmp_path, "a.md", "one\x0ctwo\nthree\u2028four\nfive\n")
    pre = SpecificationPreprocessor(tmp_path)
    result = pre.process_document(document("a.md", DocumentFormat.MD))
    locations = [n.source.location for n in result.nodes]
    assert locations == ["line:1", "line:2", "line:3"], locations


def test_diagram_text_formats_kinds(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    for fmt, kind in (
        (DocumentFormat.PLANTUML, DocumentNodeKind.DIAGRAM),
        (DocumentFormat.DOT, DocumentNodeKind.DIAGRAM),
        (DocumentFormat.SDC, DocumentNodeKind.TEXT),
        (DocumentFormat.TEX, DocumentNodeKind.TEXT),
    ):
        write(tmp_path, f"x.{fmt.value}", "line\n")
        got = pre.process_document(document(f"x.{fmt.value}", fmt))
        assert got.nodes[0].kind is kind


def test_structured_table_xml_formats(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    write(tmp_path, "a.json", '{"a": [1, 2, {"b": null}]}')
    assert pre.process_document(document("a.json", DocumentFormat.JSON)).nodes[0].content == {
        "a": [1, 2, {"b": None}]
    }
    write(tmp_path, "a.yaml", "a:\n  - 1\n  - x: y\n")
    assert pre.process_document(document("a.yaml", DocumentFormat.YAML)).nodes[0].content == {
        "a": [1, {"x": "y"}]
    }
    write(tmp_path, "a.csv", 'h1,h2\n1,"two, with comma"\n')
    csv_node = pre.process_document(document("a.csv", DocumentFormat.CSV)).nodes[0]
    assert csv_node.content == [["h1", "h2"], ["1", "two, with comma"]]
    write(tmp_path, "a.xml", '<root a="1"><child>text</child></root>')
    xml = pre.process_document(document("a.xml", DocumentFormat.XML)).nodes[0]
    assert xml.content["tag"] == "root" and xml.content["attributes"] == {"a": "1"}
    write(tmp_path, "a.svg", '<svg xmlns="http://www.w3.org/2000/svg"><rect/></svg>')
    svg = pre.process_document(document("a.svg", DocumentFormat.SVG)).nodes[0]
    assert svg.kind is DocumentNodeKind.DIAGRAM
    with pytest.raises(ValueError, match="does not exist"):
        pre.process_document(document("missing.md", DocumentFormat.MD))


def test_xml_hostile_inputs_are_rejected(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    secret = write(tmp_path, "secret.txt", "TOP-SECRET")
    xxe = f'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///{secret.as_posix()}">]><r>&x;</r>'
    write(tmp_path, "xxe.xml", xxe)
    with pytest.raises(Exception) as excinfo:
        pre.process_document(document("xxe.xml", DocumentFormat.XML))
    assert "TOP-SECRET" not in str(excinfo.value)
    bomb = (
        '<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
        '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
        '<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]>'
        "<r>&lol3;</r>"
    )
    write(tmp_path, "bomb.xml", bomb)
    with pytest.raises(Exception):
        pre.process_document(document("bomb.xml", DocumentFormat.XML))


def test_json_with_utf8_bom_is_accepted(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    (tmp_path / "bom.json").write_bytes(b"\xef\xbb\xbf" + b'{"a": 1}')
    try:
        got = pre.process_document(document("bom.json", DocumentFormat.JSON))
        outcome = ("parsed", got.nodes[0].content)
    except Exception as error:
        outcome = (type(error).__name__, str(error))
    assert outcome[0] == "parsed", outcome


def test_nested_json_within_the_depth_limit_survives_process_and_persist(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    depth = 60
    write(tmp_path, "deep.json", "[" * depth + "]" * depth)
    got = pre.process_document(document("deep.json", DocumentFormat.JSON))
    persisted = pre.persist_tree(got)
    assert yaml.safe_load(persisted.read_text("utf-8"))["document_id"] == "d1"


def test_json_nested_beyond_the_depth_limit_is_rejected_naming_the_document(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    depth = 120
    write(tmp_path, "deep.json", "[" * depth + "]" * depth)
    with pytest.raises(ValueError) as excinfo:
        pre.process_document(document("deep.json", DocumentFormat.JSON))
    assert 'Specification document "deep.json"' in str(excinfo.value)
    assert "nests more than 64 levels deep" in str(excinfo.value)


def test_csv_large_field(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    write(tmp_path, "big.csv", "h\n" + "x" * 200_000 + "\n")
    try:
        got = pre.process_document(document("big.csv", DocumentFormat.CSV))
        outcome = ("parsed", len(got.nodes[0].content[1][0]))
    except Exception as error:
        outcome = (type(error).__name__, str(error))
    assert outcome[0] == "parsed", outcome


def test_docx_xlsx_pdf_png_vsdx_formats(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    from pypdf import PdfWriter

    pre = SpecificationPreprocessor(tmp_path)
    png = tmp_path / "pic.png"
    Image.new("RGB", (8, 8), (255, 0, 0)).save(png)
    docx = Document()
    docx.add_paragraph("First paragraph")
    docx.add_paragraph("")
    docx.add_picture(str(png))
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(1, 1).text = "d"
    docx.save(tmp_path / "a.docx")
    got = pre.process_document(document("a.docx", DocumentFormat.DOCX))
    kinds = [n.kind for n in got.nodes]
    assert DocumentNodeKind.IMAGE in kinds and DocumentNodeKind.TABLE in kinds
    image_nodes = [n for n in got.nodes if n.kind is DocumentNodeKind.IMAGE]
    assert image_nodes[0].vision_status is VisionStatus.PENDING
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["h", "v"])
    workbook.active.append([1, 2.5])
    workbook.create_sheet("Second").append(["x"])
    workbook.save(tmp_path / "a.xlsx")
    sheets = pre.process_document(document("a.xlsx", DocumentFormat.XLSX)).nodes
    assert [n.node_id for n in sheets] == ["sheet-Sheet1", "sheet-Second"]
    assert sheets[0].content == [["h", "v"], [1, 2.5]]
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with (tmp_path / "a.pdf").open("wb") as handle:
        writer.write(handle)
    pdf = pre.process_document(document("a.pdf", DocumentFormat.PDF))
    assert pdf.nodes[0].node_id == "page-1-text"
    png_tree = pre.process_document(document("pic.png", DocumentFormat.PNG))
    assert png_tree.nodes[0].kind is DocumentNodeKind.IMAGE
    with zipfile.ZipFile(tmp_path / "a.vsdx", "w") as package:
        package.writestr("visio/pages/page1.xml", "<Page/>")
        package.writestr("visio/pages/page2.xml", "<Page2/>")
    vsdx = pre.process_document(document("a.vsdx", DocumentFormat.VSDX))
    assert [n.node_id for n in vsdx.nodes] == ["vsdx-1", "vsdx-2"]
    with zipfile.ZipFile(tmp_path / "empty.vsdx", "w") as package:
        package.writestr("other.txt", "x")
    with pytest.raises(ValueError, match="no Visio page XML parts"):
        pre.process_document(document("empty.vsdx", DocumentFormat.VSDX))


def test_corrupt_binary_documents_fail_with_exceptions_not_hangs(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    for name, fmt in (
        ("bad.docx", DocumentFormat.DOCX),
        ("bad.pdf", DocumentFormat.PDF),
        ("bad.xlsx", DocumentFormat.XLSX),
        ("bad.vsdx", DocumentFormat.VSDX),
    ):
        (tmp_path / name).write_bytes(b"not really a document")
        with pytest.raises(Exception):
            pre.process_document(document(name, fmt))


def test_vsdx_page_parts_over_the_decompression_limit_are_rejected(tmp_path):
    package = tmp_path / "bomb.vsdx"
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as handle:
        handle.writestr("visio/pages/page1.xml", b"0" * 120_000_000)
    pre = SpecificationPreprocessor(tmp_path)
    with pytest.raises(ValueError, match=r"visio/pages/page1\.xml.*decompress"):
        pre.process_document(document("bomb.vsdx", DocumentFormat.VSDX))


def test_resolve_images_rules(tmp_path):
    img = node(
        "image-1",
        {"filename": "p.png"},
        kind=DocumentNodeKind.IMAGE,
        vision_status=VisionStatus.PENDING,
    )
    text = node("line-1", "text")
    base = tree(nodes=[img, text])
    pre = SpecificationPreprocessor(tmp_path)
    good = ScriptedVisionAdapter([VisionProposal(confidence=0.8, structure={"signals": ["clk"]})])
    resolved = arun(pre.resolve_images(base, good))
    assert resolved.nodes[0].vision_status is VisionStatus.ACCEPTED
    assert resolved.nodes[0].resolved_structure == {"signals": ["clk"]}
    assert resolved.nodes[1] == text
    empty_structure = ScriptedVisionAdapter([VisionProposal(confidence=1.0, structure={})] * 3)
    review = arun(pre.resolve_images(base, empty_structure))
    assert review.nodes[0].vision_status is VisionStatus.REVIEW_REQUIRED
    below = ScriptedVisionAdapter([VisionProposal(confidence=0.79, structure={"a": 1})] * 3)
    assert (
        arun(pre.resolve_images(base, below)).nodes[0].vision_status is VisionStatus.REVIEW_REQUIRED
    )
    retry = ScriptedVisionAdapter(
        [
            VisionProposal(confidence=0.1, structure={"a": 1}),
            VisionProposal(confidence=0.9, structure={"b": 2}),
        ]
    )
    assert arun(pre.resolve_images(base, retry)).nodes[0].resolved_structure == {"b": 2}
    calls = []

    class Counting:
        async def extract(self, n):
            calls.append(1)
            return VisionProposal(confidence=0.0, structure={})

    arun(pre.resolve_images(base, Counting(), max_attempts=5))
    assert len(calls) == 5
    unconfigured = arun(pre.resolve_images(base, UnconfiguredVisionAdapter()))
    assert unconfigured.nodes[0].vision_status is VisionStatus.REVIEW_REQUIRED
    assert unconfigured.nodes[0].resolved_structure is None

    class Raising:
        async def extract(self, n):
            raise RuntimeError("vision backend down")

    with pytest.raises(RuntimeError, match="vision backend down"):
        arun(pre.resolve_images(base, Raising()))


def test_persist_tree_roundtrip(tmp_path):
    write(tmp_path, "a.md", "alpha\nbeta\n")
    pre = SpecificationPreprocessor(tmp_path)
    got = pre.process_document(document("a.md", DocumentFormat.MD))
    target = pre.persist_tree(got)
    assert target == tmp_path / "processed" / "d1.document.yaml"
    loaded = yaml.safe_load(target.read_text("utf-8"))
    assert loaded == json.loads(json.dumps(got.model_dump(mode="json")))


def test_vision_results_are_consumed_downstream(tmp_path):
    from nailong_agent_sdk.specifications.retrieval import specification_retrieval_documents

    img = node(
        "image-1",
        {"filename": "block.png"},
        kind=DocumentNodeKind.IMAGE,
        vision_status=VisionStatus.ACCEPTED,
        resolved_structure={"signals": ["rst_n_unique_token"]},
    )
    documents = specification_retrieval_documents("snap", [tree(nodes=[img])])
    assert any("rst_n_unique_token" in d.text for d in documents)


def test_unreviewed_image_nodes_are_flagged_by_gate1():
    img = node(
        "image-1",
        {"filename": "p.png"},
        kind=DocumentNodeKind.IMAGE,
        vision_status=VisionStatus.REVIEW_REQUIRED,
    )
    specification = spec(reqs=[req("R1")], trees=[tree(nodes=[img])])
    _, report = SpecificationGate().validate(specification, required_categories=set())
    assert report.gaps, "REVIEW_REQUIRED image content passes Gate 1 with no gap"


def test_keywords_from_task_ascii_only():
    assert keywords_from_task("Verify REQ-FUNC-001 and clk_div") == {
        "verify",
        "req-func-001",
        "and",
        "clk_div",
    }


def test_gate_basic_checks_and_severities():
    requirements = [
        req("R1", deps=["R2"]),
        req("R2", deps=["R9"], checks=()),
        req("R3", text="   "),
        req("R4", deps=["R4"]),
    ]
    specification = spec(reqs=requirements, trees=[tree()])
    graph, report = SpecificationGate().validate(
        specification,
        required_categories={SpecificationCategory.FUNCTIONAL, SpecificationCategory.PPA},
    )
    by_source = {}
    for gap in report.gaps:
        by_source.setdefault(gap.source, []).append(gap)
    absence = [g for g in report.gaps if g.type is GapType.ABSENCE]
    assert any(
        g.locations == ["category:ppa"] and g.severity is GapSeverity.CRITICAL for g in absence
    )
    assert any(g.locations == ["R3"] for g in absence)
    missing = [g for g in report.gaps if g.locations == ["R2", "R9"]]
    assert len(missing) == 1 and missing[0].blast_radius == 2
    verifiability = [g for g in report.gaps if g.type is GapType.VERIFIABILITY]
    assert [g.locations for g in verifiability] == [["R2"]]
    cycles = [g for g in report.gaps if g.description.startswith("Requirement dependency graph")]
    assert [g.locations for g in cycles] == [["R4", "R4"]]
    assert graph.nodes == ["R1", "R2", "R3", "R4"]
    assert report.summary()["critical"] >= 3


def _brute_cycle_exists(nodes, edges):
    adjacency = {n: [] for n in nodes}
    for a, b in edges:
        adjacency[a].append(b)
    state = {}

    def visit(n):
        state[n] = 1
        for m in adjacency[n]:
            if state.get(m) == 1:
                return True
            if m not in state and visit(m):
                return True
        state[n] = 2
        return False

    return any(n not in state and visit(n) for n in nodes)


def test_cycle_detection_matches_brute_force_on_random_graphs():
    from nailong_agent_sdk.foundations.dependency_graph import deterministic_cycles

    rng = random.Random(7)
    checked = 0
    for _ in range(400):
        n = rng.randint(1, 8)
        nodes = [f"N{i}" for i in range(n)]
        edges = [(rng.choice(nodes), rng.choice(nodes)) for _ in range(rng.randint(0, 12))]
        cycles = deterministic_cycles(nodes, edges)
        assert bool(cycles) == _brute_cycle_exists(nodes, edges), (nodes, edges)
        edge_set = set(edges)
        for cycle in cycles:
            assert cycle[0] == cycle[-1]
            assert all((a, b) in edge_set for a, b in itertools.pairwise(cycle))
        checked += 1
    assert checked == 400


def test_gate_scales_to_large_dependency_graphs():
    n = 20_000
    chain = [req(f"R{i}", deps=[f"R{i + 1}"] if i + 1 < n else []) for i in range(n)]
    started = time.monotonic()
    _, report = SpecificationGate().validate(
        spec(reqs=chain, trees=[tree()]), required_categories=set()
    )
    time.monotonic() - started
    assert report.gaps == []
    ring_n = 5_000
    ring = [req(f"R{i}", deps=[f"R{(i + 1) % ring_n}"]) for i in range(ring_n)]
    started = time.monotonic()
    _, ring_report = SpecificationGate().validate(
        spec(reqs=ring, trees=[tree()]), required_categories=set()
    )
    time.monotonic() - started
    assert len(ring_report.gaps) == 1 and len(ring_report.gaps[0].locations) == ring_n + 1


def test_duplicate_dependency_entries_do_not_duplicate_gaps():
    requirements = [req("R1", deps=["R2", "R2"]), req("R2", deps=["R1"])]
    _, report = SpecificationGate().validate(
        spec(reqs=requirements, trees=[tree()]), required_categories=set()
    )
    cycle_gaps = [
        g for g in report.gaps if g.description.startswith("Requirement dependency graph")
    ]
    assert len(cycle_gaps) == 1


def _finding(**overrides):
    values = dict(
        finding_id="F1",
        type=GapType.AMBIGUITY,
        requirement_ids=["R1"],
        source_refs=[sref(loc="line:R1")],
        description="ambiguous",
        suggested_fix="clarify",
        analysis_provider="prov",
        analysis_receipt_digest="a" * 64,
    )
    values.update(overrides)
    return SemanticGapFinding(**values)


def test_semantic_finding_admission_rules_and_model_validation():
    specification = spec(reqs=[req("R1"), req("R2", deps=["R1"])], trees=[tree()])
    gate = SpecificationGate()
    graph, report = gate.validate(
        specification,
        required_categories=set(),
        semantic_findings=[
            _finding(finding_id="B-ok"),
            _finding(finding_id="A-unknown", requirement_ids=["R404"]),
            _finding(finding_id="C-foreign-source", source_refs=[sref(loc="line:ZZ")]),
            _finding(finding_id="D-dup"),
            _finding(finding_id="D-dup"),
            _finding(
                finding_id="E-incons",
                type=GapType.INCONSISTENCY,
                requirement_ids=["R1", "R2"],
                source_refs=[sref(loc="line:R1"), sref(loc="line:R2")],
            ),
        ],
    )
    admissions = {a.finding_id: a for a in report.semantic_admissions}
    assert admissions["B-ok"].accepted and admissions["E-incons"].accepted
    assert not admissions["A-unknown"].accepted and "R404" in admissions["A-unknown"].reason
    assert not admissions["C-foreign-source"].accepted
    assert [a.accepted for a in report.semantic_admissions if a.finding_id == "D-dup"] == [
        True,
        False,
    ]
    repeated = [a for a in report.semantic_admissions if a.finding_id == "D-dup"][1]
    assert repeated.reason == (
        'Semantic finding ID "D-dup" repeats an earlier finding; '
        "finding IDs must be unique per Gate 1 evaluation."
    )
    accepted_gap = next(g for g in report.gaps if g.source.startswith("semantic-analysis:prov"))
    assert accepted_gap.model_analysis_required is True and accepted_gap.blast_radius >= 1
    with pytest.raises(ValueError, match="CRITICAL"):
        _finding(severity=GapSeverity.CRITICAL)
    with pytest.raises(ValueError, match="semantic GapType"):
        _finding(type=GapType.ABSENCE)
    with pytest.raises(ValueError, match="at least two"):
        _finding(type=GapType.INCONSISTENCY)
    with pytest.raises(ValueError):
        _finding(analysis_receipt_digest="G" * 64)
    with pytest.raises(ValueError, match="unique"):
        _finding(requirement_ids=["R1", "R1"])
    assert report.summary()["semantic_findings_rejected"] == 3


def test_soft_lock_decisions():
    specification = spec(reqs=[req("R1", checks=())], trees=[tree()])
    gate = SpecificationGate()
    _, report = gate.validate(specification, required_categories=set())
    metadata = VersionMetadata(
        version="1.0.0",
        change_kind=VersionChangeKind.MAJOR,
        unified_specification_hash="not-the-real-hash",
    )
    refused = gate.soft_lock(specification, report, metadata, user_approved=False)
    assert not refused.accepted and "Designer approval" in refused.warnings[0]
    gaps = gate.soft_lock(specification, report, metadata, user_approved=True)
    assert not gaps.accepted and "proceed_with_gaps" in gaps.warnings[-1]
    forced = gate.soft_lock(
        specification, report, metadata, user_approved=True, proceed_with_gaps=True
    )
    assert (
        forced.accepted and forced.metadata.soft_locked and forced.metadata.user_override_with_gaps
    )
    clean_spec = spec(reqs=[req("R1")], trees=[tree()])
    _, clean_report = gate.validate(clean_spec, required_categories=set())
    clean = gate.soft_lock(clean_spec, clean_report, metadata, user_approved=True)
    assert clean.accepted and not clean.metadata.user_override_with_gaps


def test_soft_lock_refuses_a_report_or_metadata_written_for_another_version():
    specification = spec(version="2.0.0", reqs=[req("R1")], trees=[tree()])
    gate = SpecificationGate()
    _, report = gate.validate(specification, required_categories=set())
    metadata = VersionMetadata(
        version="2.0.0", change_kind=VersionChangeKind.MAJOR, unified_specification_hash="h"
    )
    stale_report = report.model_copy(update={"document_version": "1.9.0"})
    for approved in (True, False):
        with pytest.raises(
            ValueError,
            match=r'Gap report document_version "1.9.0" does not match the specification '
            r'version "2.0.0"',
        ):
            gate.soft_lock(specification, stale_report, metadata, user_approved=approved)
    other = metadata.model_copy(update={"version": "2.0.1"})
    with pytest.raises(
        ValueError,
        match=r'Version metadata version "2.0.1" does not match the specification '
        r'version "2.0.0"',
    ):
        gate.soft_lock(specification, report, other, user_approved=True)
    assert gate.soft_lock(specification, report, metadata, user_approved=True).accepted


def test_classify_version_change_matrix():
    base = spec(reqs=[req("R1"), req("R2")], trees=[tree()])
    assert classify_version_change(None, base)[0] is VersionChangeKind.MAJOR
    removed = spec(reqs=[req("R1")], trees=[tree()])
    assert classify_version_change(base, removed)[0] is VersionChangeKind.MAJOR
    changed = spec(reqs=[req("R1", text="different"), req("R2")], trees=[tree()])
    kind, why = classify_version_change(base, changed)
    assert kind is VersionChangeKind.MAJOR and "R1 (text)" in why[0]
    added = spec(reqs=[req("R1"), req("R2"), req("R3")], trees=[tree()])
    assert classify_version_change(base, added)[0] is VersionChangeKind.MINOR
    checks_only = spec(reqs=[req("R1", checks=("a", "b")), req("R2")], trees=[tree()])
    assert classify_version_change(base, checks_only)[0] is VersionChangeKind.PATCH
    assert classify_version_change(base, base)[0] is VersionChangeKind.PATCH
    both = spec(reqs=[req("R1", text="x"), req("R2"), req("R3")], trees=[tree()])
    assert classify_version_change(base, both)[0] is VersionChangeKind.MAJOR


def _persist(store, specification, plans=()):
    gate = SpecificationGate()
    graph, report = gate.validate(specification, required_categories=set())
    metadata = VersionMetadata(
        version=specification.version,
        change_kind=VersionChangeKind.MAJOR,
        unified_specification_hash="h",
    )
    store.persist(specification, graph, report, metadata, list(plans))


def test_gate1_store_persist_and_yaml_roundtrip_of_hostile_strings(tmp_path):
    tricky = [
        "yes",
        "no",
        "null",
        "~",
        "1e3",
        "0x1F",
        "2026-01-01",
        "  leading",
        "trailing  ",
        "tab\there",
        "line1\nline2",
        "colon: value",
        "- dash",
        "# hash",
        '"quoted"',
        "emoji \U0001f600",
        "\u2028sep",
        "\x85nel",
        "a" * 600,
        "ünï文",
        "{braces}",
        "[brackets]",
        "&anchor",
        "*alias",
        "!tag",
        "%dir",
        "@at",
        "`tick`",
        "'single'",
        "\ufeffbom",
        "\x00nul",
        "\x1bescape",
    ]
    store = Gate1ArtifactStore(tmp_path / "specifications")
    mismatches = []
    for text in tricky:
        specification = spec(reqs=[req("R1", text=text)], trees=[tree()])
        try:
            _persist(store, specification)
            loaded = yaml.safe_load(
                (tmp_path / "specifications" / "unified-specification.yaml").read_text("utf-8")
            )
            if loaded["requirements"][0]["text"] != text:
                mismatches.append((text, loaded["requirements"][0]["text"]))
        except Exception as error:
            mismatches.append((text, f"{type(error).__name__}: {str(error)[:80]}"))
    assert mismatches == [], mismatches


def test_gate1_store_leaves_no_stale_plans(tmp_path):
    store = Gate1ArtifactStore(tmp_path / "specifications")
    specification = spec(reqs=[req("R1")], trees=[tree()])
    _persist(store, specification, plans=[{"n": 1}, {"n": 2}, {"n": 3}])
    _persist(store, specification, plans=[{"n": 10}])
    remaining = sorted(p.name for p in (tmp_path / "specifications" / "plans").iterdir())
    assert remaining == ["plan-1.yaml"], remaining


def _egraph(node_specs, relations=()):
    nodes = [
        EvidenceNode(
            node_id=node_id,
            kind=kind,
            content=content,
            source=sref(loc=f"line:{node_id}"),
            token_cost=cost,
            authority_tier=authority,
            aliases=aliases,
        )
        for node_id, kind, content, cost, authority, aliases in node_specs
    ]
    rels = [
        EvidenceRelation(from_node_id=a, to_node_id=b, kind=kind, source=sref(loc=f"{a}->{b}"))
        for a, b, kind in relations
    ]
    return EvidenceGraph(snapshot_id="snap", nodes=nodes, relations=rels)


def test_evidence_closure_aliases_budget_and_witnesses():
    K = EvidenceNodeKind
    graph = _egraph(
        [
            ("T", K.REQUIREMENT, "target", 10, 0, ["alias-t"]),
            ("A", K.INTERFACE, "iface", 10, 0, []),
            ("B", K.SIGNAL, "sig", 10, 0, []),
            ("C", K.DECISION, "optional note about target", 5, 0, []),
            ("D", K.SOURCE_SPAN, "advisory only", 5, 0, []),
        ],
        [
            ("T", "A", EvidenceRelationKind.REQUIRES),
            ("A", "B", EvidenceRelationKind.INTERFACE_CONSTRAINT),
            ("T", "D", EvidenceRelationKind.ADVISORY),
        ],
    )
    selector = StructuralContextSelector()
    policy = EvidenceSelectionPolicy(policy_id="p")
    result = selector.select(
        graph,
        EvidenceSelectionRequest(snapshot_id="snap", target_ids=["alias-t"], token_budget=100),
        policy,
    )
    assert result.status is EvidenceSelectionStatus.SELECTED
    assert result.mandatory_node_ids == ["A", "B", "T"]
    assert {n.node_id for n in result.selected_nodes} == {"T", "A", "B", "C", "D"}
    witness = next(w for w in result.witnesses if w.node_id == "B")
    assert [r.to_node_id for r in witness.relation_path] == [
        "A",
        "B",
    ] and witness.root_target_id == "T"
    tight = selector.select(
        graph,
        EvidenceSelectionRequest(snapshot_id="snap", target_ids=["T"], token_budget=29),
        policy,
    )
    assert (
        tight.status is EvidenceSelectionStatus.INFEASIBLE_REQUIRED_CLOSURE
        and tight.selected_nodes == []
    )
    with pytest.raises(ValueError, match='"nope" is unknown'):
        selector.select(
            graph,
            EvidenceSelectionRequest(snapshot_id="snap", target_ids=["nope"], token_budget=100),
            policy,
        )
    with pytest.raises(ValueError, match="unexpected graph snapshot"):
        selector.select(
            graph,
            EvidenceSelectionRequest(snapshot_id="other", target_ids=["T"], token_budget=100),
            policy,
        )
    ambiguous = _egraph(
        [("X", K.REQUIREMENT, "x", 1, 0, ["dup"]), ("Y", K.REQUIREMENT, "y", 1, 0, ["dup"])]
    )
    with pytest.raises(ValueError, match="ambiguous"):
        selector.select(
            ambiguous,
            EvidenceSelectionRequest(snapshot_id="snap", target_ids=["dup"], token_budget=10),
            policy,
        )


def test_evidence_strict_flag_has_an_effect():
    K = EvidenceNodeKind
    graph = _egraph([("T", K.REQUIREMENT, "t", 1, 0, [])])
    selector = StructuralContextSelector()
    policy = EvidenceSelectionPolicy(policy_id="p")
    try:
        selector.select(
            graph,
            EvidenceSelectionRequest(
                snapshot_id="snap", target_ids=["T", "ghost"], token_budget=10, strict=False
            ),
            policy,
        )
        outcome = "tolerated"
    except ValueError as error:
        outcome = f"ValueError: {error}"
    assert outcome == "tolerated", outcome
    result = selector.select(
        graph,
        EvidenceSelectionRequest(
            snapshot_id="snap", target_ids=["T", "ghost"], token_budget=10, strict=False
        ),
        policy,
    )
    assert 'Evidence target "ghost" is unknown (ignored).' in result.diagnostics
    assert [node.node_id for node in result.selected_nodes] == ["T"]
    with pytest.raises(ValueError, match='Evidence target "ghost" is unknown'):
        selector.select(
            graph,
            EvidenceSelectionRequest(
                snapshot_id="snap", target_ids=["T", "ghost"], token_budget=10
            ),
            policy,
        )
    with pytest.raises(ValueError, match="no resolvable target"):
        selector.select(
            graph,
            EvidenceSelectionRequest(
                snapshot_id="snap", target_ids=["ghost"], token_budget=10, strict=False
            ),
            policy,
        )


def test_evidence_optional_packing_is_optimal_against_brute_force():
    K = EvidenceNodeKind
    rng = random.Random(11)
    for trial in range(40):
        count = rng.randint(1, 9)
        specs = [("T", K.REQUIREMENT, "target text", 3, 0, [])]
        for i in range(count):
            specs.append(
                (
                    f"O{i}",
                    rng.choice(list(K)),
                    "word " * rng.randint(0, 3),
                    rng.randint(1, 9),
                    rng.randint(0, 5),
                    [],
                )
            )
        graph = _egraph(specs)
        budget = rng.randint(3, 40)
        policy = EvidenceSelectionPolicy(policy_id="p")
        result = StructuralContextSelector().select(
            graph,
            EvidenceSelectionRequest(
                snapshot_id="snap", target_ids=["T"], token_budget=budget, task_text="word target"
            ),
            policy,
        )
        nodes = {n.node_id: n for n in graph.nodes}
        task_terms = {"word", "target"}
        utilities = {}
        for node_id, n in nodes.items():
            if node_id == "T":
                continue
            content_terms = {t for t in str(n.content).lower().split() if len(t) > 1}
            utilities[node_id] = (
                policy.node_kind_utility[n.kind]
                + n.authority_tier * policy.authority_multiplier
                + len(task_terms & content_terms) * policy.lexical_overlap_weight
            )
        remaining = budget - 3
        best = 0
        ids = list(utilities)
        for mask in range(1 << len(ids)):
            cost = sum(nodes[ids[i]].token_cost for i in range(len(ids)) if mask >> i & 1)
            if cost <= remaining:
                best = max(best, sum(utilities[ids[i]] for i in range(len(ids)) if mask >> i & 1))
        got = sum(utilities[i] for i in result.selected_optional_node_ids)
        assert got == best, (trial, got, best)
        assert result.token_cost <= budget


def test_evidence_large_budget_behaviour():
    K = EvidenceNodeKind
    rng = random.Random(5)
    specs = [("T", K.REQUIREMENT, "target", 10, 0, [])]
    for i in range(300):
        specs.append(
            (
                f"O{i}",
                rng.choice(list(K)),
                f"word{i % 7}",
                rng.randint(50, 500),
                rng.randint(0, 20),
                [],
            )
        )
    graph = _egraph(specs)
    policy = EvidenceSelectionPolicy(policy_id="p")
    report = {}
    for budget in (20_000, 50_000, 100_000):
        started = time.monotonic()
        result = StructuralContextSelector().select(
            graph,
            EvidenceSelectionRequest(
                snapshot_id="snap", target_ids=["T"], token_budget=budget, task_text="word1"
            ),
            policy,
        )
        report[budget] = {
            "seconds": round(time.monotonic() - started, 2),
            "status": result.status.value,
            "cost": result.token_cost,
        }
    assert all(entry["status"] == "selected" for entry in report.values()), report
    assert max(entry["seconds"] for entry in report.values()) < 15, report
