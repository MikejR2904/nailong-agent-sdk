import zipfile

import pytest
import yaml

from nailong_agent_sdk.specifications.documents import (
    DocumentFormat,
    DocumentNodeKind,
    VisionStatus,
)
from nailong_agent_sdk.specifications.preprocessing import SpecificationPreprocessor
from tests.support.document_cases import document, write


def test_diagram_text_formats_kinds(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    for fmt, kind in (
        (DocumentFormat.PLANTUML, DocumentNodeKind.DIAGRAM),
        (DocumentFormat.DOT, DocumentNodeKind.DIAGRAM),
        (DocumentFormat.MERMAID, DocumentNodeKind.DIAGRAM),
        (DocumentFormat.TEX, DocumentNodeKind.TEXT),
    ):
        write(tmp_path, f"x.{fmt.value}", "line\n")
        got = pre.process_document(document(f"x.{fmt.value}", fmt))
        assert got.nodes[0].kind is kind


def test_structured_table_xml_formats(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    write(tmp_path, "a.json", '{"a": [1, 2, {"b": null}]}')
    assert pre.process_document(document("a.json", DocumentFormat.JSON)).nodes[0].content == {
        "a": [1, 2, {"b": None}]
    }
    write(tmp_path, "a.yaml", "a:\n  - 1\n  - x: y\n")
    assert pre.process_document(document("a.yaml", DocumentFormat.YAML)).nodes[0].content == {
        "a": [1, {"x": "y"}]
    }
    write(tmp_path, "a.csv", 'h1,h2\n1,"two, with comma"\n')
    csv_node = pre.process_document(document("a.csv", DocumentFormat.CSV)).nodes[0]
    assert csv_node.content == [["h1", "h2"], ["1", "two, with comma"]]
    write(tmp_path, "a.xml", '<root a="1"><child>text</child></root>')
    xml = pre.process_document(document("a.xml", DocumentFormat.XML)).nodes[0]
    assert xml.content["tag"] == "root" and xml.content["attributes"] == {"a": "1"}
    write(tmp_path, "a.svg", '<svg xmlns="http://www.w3.org/2000/svg"><rect/></svg>')
    svg = pre.process_document(document("a.svg", DocumentFormat.SVG)).nodes[0]
    assert svg.kind is DocumentNodeKind.DIAGRAM
    with pytest.raises(ValueError, match="does not exist"):
        pre.process_document(document("missing.md", DocumentFormat.MD))


def test_xml_hostile_inputs_are_rejected(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    secret = write(tmp_path, "secret.txt", "TOP-SECRET")
    xxe = f'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///{secret.as_posix()}">]><r>&x;</r>'
    write(tmp_path, "xxe.xml", xxe)
    with pytest.raises(Exception) as excinfo:
        pre.process_document(document("xxe.xml", DocumentFormat.XML))
    assert "TOP-SECRET" not in str(excinfo.value)
    bomb = (
        '<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
        '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
        '<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]>'
        "<r>&lol3;</r>"
    )
    write(tmp_path, "bomb.xml", bomb)
    with pytest.raises(Exception):
        pre.process_document(document("bomb.xml", DocumentFormat.XML))


def test_json_with_utf8_bom_is_accepted(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    (tmp_path / "bom.json").write_bytes(b"\xef\xbb\xbf" + b'{"a": 1}')
    try:
        got = pre.process_document(document("bom.json", DocumentFormat.JSON))
        outcome = ("parsed", got.nodes[0].content)
    except Exception as error:
        outcome = (type(error).__name__, str(error))
    assert outcome[0] == "parsed", outcome


def test_nested_json_within_the_depth_limit_survives_process_and_persist(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    depth = 60
    write(tmp_path, "deep.json", "[" * depth + "]" * depth)
    got = pre.process_document(document("deep.json", DocumentFormat.JSON))
    persisted = pre.persist_tree(got)
    assert yaml.safe_load(persisted.read_text("utf-8"))["document_id"] == "d1"


def test_json_nested_beyond_the_depth_limit_is_rejected_naming_the_document(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    depth = 120
    write(tmp_path, "deep.json", "[" * depth + "]" * depth)
    with pytest.raises(ValueError) as excinfo:
        pre.process_document(document("deep.json", DocumentFormat.JSON))
    assert 'Specification document "deep.json"' in str(excinfo.value)
    assert "nests more than 64 levels deep" in str(excinfo.value)


def test_csv_large_field(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    write(tmp_path, "big.csv", "h\n" + "x" * 200_000 + "\n")
    try:
        got = pre.process_document(document("big.csv", DocumentFormat.CSV))
        outcome = ("parsed", len(got.nodes[0].content[1][0]))
    except Exception as error:
        outcome = (type(error).__name__, str(error))
    assert outcome[0] == "parsed", outcome


def test_docx_xlsx_pdf_png_vsdx_formats(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    from pypdf import PdfWriter

    pre = SpecificationPreprocessor(tmp_path)
    png = tmp_path / "pic.png"
    Image.new("RGB", (8, 8), (255, 0, 0)).save(png)
    docx = Document()
    docx.add_paragraph("First paragraph")
    docx.add_paragraph("")
    docx.add_picture(str(png))
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(1, 1).text = "d"
    docx.save(tmp_path / "a.docx")
    got = pre.process_document(document("a.docx", DocumentFormat.DOCX))
    kinds = [n.kind for n in got.nodes]
    assert DocumentNodeKind.IMAGE in kinds and DocumentNodeKind.TABLE in kinds
    image_nodes = [n for n in got.nodes if n.kind is DocumentNodeKind.IMAGE]
    assert image_nodes[0].vision_status is VisionStatus.PENDING
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["h", "v"])
    workbook.active.append([1, 2.5])
    workbook.create_sheet("Second").append(["x"])
    workbook.save(tmp_path / "a.xlsx")
    sheets = pre.process_document(document("a.xlsx", DocumentFormat.XLSX)).nodes
    assert [n.node_id for n in sheets] == ["sheet-Sheet1", "sheet-Second"]
    assert sheets[0].content == [["h", "v"], [1, 2.5]]
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with (tmp_path / "a.pdf").open("wb") as handle:
        writer.write(handle)
    pdf = pre.process_document(document("a.pdf", DocumentFormat.PDF))
    assert pdf.nodes[0].node_id == "page-1-text"
    png_tree = pre.process_document(document("pic.png", DocumentFormat.PNG))
    assert png_tree.nodes[0].kind is DocumentNodeKind.IMAGE
    with zipfile.ZipFile(tmp_path / "a.vsdx", "w") as package:
        package.writestr("visio/pages/page1.xml", "<Page/>")
        package.writestr("visio/pages/page2.xml", "<Page2/>")
    vsdx = pre.process_document(document("a.vsdx", DocumentFormat.VSDX))
    assert [n.node_id for n in vsdx.nodes] == ["vsdx-1", "vsdx-2"]
    with zipfile.ZipFile(tmp_path / "empty.vsdx", "w") as package:
        package.writestr("other.txt", "x")
    with pytest.raises(ValueError, match="no Visio page XML parts"):
        pre.process_document(document("empty.vsdx", DocumentFormat.VSDX))


def test_corrupt_binary_documents_fail_with_exceptions_not_hangs(tmp_path):
    pre = SpecificationPreprocessor(tmp_path)
    for name, fmt in (
        ("bad.docx", DocumentFormat.DOCX),
        ("bad.pdf", DocumentFormat.PDF),
        ("bad.xlsx", DocumentFormat.XLSX),
        ("bad.vsdx", DocumentFormat.VSDX),
    ):
        (tmp_path / name).write_bytes(b"not really a document")
        with pytest.raises(Exception):
            pre.process_document(document(name, fmt))


def test_vsdx_page_parts_over_the_decompression_limit_are_rejected(tmp_path):
    package = tmp_path / "bomb.vsdx"
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as handle:
        handle.writestr("visio/pages/page1.xml", b"0" * 120_000_000)
    pre = SpecificationPreprocessor(tmp_path)
    with pytest.raises(ValueError, match=r"visio/pages/page1\.xml.*decompress"):
        pre.process_document(document("bomb.vsdx", DocumentFormat.VSDX))
